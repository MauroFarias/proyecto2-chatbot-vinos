"""Copia las respuestas del LLM de resultados.json a la tabla de la sección 4 del informe.

Uso (desde la raíz del proyecto, después de correr la evaluación):
    python -m evaluacion.evaluar
    python -m evaluacion.rellenar_informe

Reemplaza los [COMPLETAR] de la columna "Respuesta LLM" en informe/informe.md y
informe/informe.docx, y muestra el tiempo promedio de cada motor. NO completa la
clasificación (Correcta/Parcial/Incorrecta), el análisis ni las conclusiones:
eso requiere leer las respuestas.
"""
import json
from pathlib import Path

from docx import Document

RAIZ = Path(__file__).resolve().parent.parent
RESULTADOS = RAIZ / "evaluacion" / "resultados.json"
INFORME_MD = RAIZ / "informe" / "informe.md"
INFORME_DOCX = RAIZ / "informe" / "informe.docx"
MARCA = "[COMPLETAR]"
ENCABEZADO = "Respuesta LLM"


def cargar_respuestas_llm() -> dict[int, str]:
    """Retorna {número de pregunta: texto de la respuesta del LLM}."""
    resultados = json.loads(RESULTADOS.read_text(encoding="utf-8"))
    if not all("llm" in fila for fila in resultados):
        raise SystemExit("resultados.json no tiene respuestas del LLM. Corre primero la evaluación.")
    return {fila["n"]: fila["llm"]["texto"].strip() for fila in resultados}


def rellenar_md(respuestas: dict[int, str]) -> int:
    """Rellena la tabla cuyo encabezado contiene 'Respuesta LLM'. Retorna cuántas celdas cambió."""
    lineas = INFORME_MD.read_text(encoding="utf-8").splitlines()
    en_tabla, cambios = False, 0
    for i, linea in enumerate(lineas):
        if linea.startswith("|") and ENCABEZADO in linea:
            en_tabla = True
        elif en_tabla and not linea.startswith("|"):
            en_tabla = False
        elif en_tabla and linea.endswith(f"| {MARCA} |"):
            numero = int(linea.split("|")[1])
            texto = respuestas[numero].replace("\n", " ").replace("|", "/")
            lineas[i] = linea[: -len(MARCA) - 2] + f"{texto} |"
            cambios += 1
    INFORME_MD.write_text("\n".join(lineas) + "\n", encoding="utf-8")
    return cambios


def rellenar_docx(respuestas: dict[int, str]) -> int:
    """Rellena la columna 'Respuesta LLM' del .docx y quita el resaltado amarillo."""
    documento = Document(INFORME_DOCX)
    cambios = 0
    for tabla in documento.tables:
        if ENCABEZADO not in tabla.rows[0].cells[-1].text:
            continue
        for fila in tabla.rows[1:]:
            celda = fila.cells[-1]
            if MARCA not in celda.text:
                continue
            numero = int(fila.cells[0].text)
            parrafo = celda.paragraphs[0]
            base = parrafo.runs[0]
            base.text = respuestas[numero]
            base.font.highlight_color = None
            for extra in parrafo.runs[1:]:
                extra._element.getparent().remove(extra._element)
            cambios += 1
    documento.save(INFORME_DOCX)
    return cambios


def mostrar_tiempos() -> None:
    resultados = json.loads(RESULTADOS.read_text(encoding="utf-8"))
    for motor in ("prolog", "llm"):
        tiempos = [f[motor]["segundos"] for f in resultados if motor in f]
        if tiempos:
            print(f"Tiempo promedio {motor}: {sum(tiempos) / len(tiempos):.2f} s")


def main() -> None:
    respuestas = cargar_respuestas_llm()
    print(f"informe.md:   {rellenar_md(respuestas)} respuestas copiadas")
    print(f"informe.docx: {rellenar_docx(respuestas)} respuestas copiadas")
    mostrar_tiempos()
    print("Quedan por completar a mano: clasificación 5.2, tiempo del LLM, 5.4 y conclusiones.")


if __name__ == "__main__":
    main()
